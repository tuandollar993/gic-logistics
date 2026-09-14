import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.models import Lot, OperatingCost, Target
from app.services.calculator import CalculatorService

class ExcelExporterService:
    @staticmethod
    def export_monthly_report(month: int, year: int) -> io.BytesIO:
        wb = openpyxl.Workbook()
        
        font_title = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        font_subtitle = Font(name="Arial", size=9.5, italic=True, color="475569")
        font_header = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
        font_kpi_title = Font(name="Arial", size=8.5, bold=True, color="475569")
        font_kpi_val = Font(name="Arial", size=11, bold=True, color="0F172A")
        font_data = Font(name="Arial", size=9.5, color="1E293B")
        font_total = Font(name="Arial", size=10, bold=True, color="0F172A")
        
        fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        thin_side = Side(style='thin', color='CBD5E1')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Side(style='double', color='0F172A')
        total_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)
        
        num_fmt_currency = '#,##0'
        num_fmt_percent = '0.0%'
        
        lots = Lot.sales_lots_query().filter_by(month=month, year=year).order_by(Lot.id.asc()).all()
        unresolved_groups = Lot.unresolved_cpvh_query().filter_by(month=month, year=year).all()
        unresolved_cost_total = sum(u.total_operating_cost for u in unresolved_groups)
        kpi = CalculatorService.get_monthly_kpi(month, year)
        
        # SHEET 1: BÁO CÁO KINH DOANH
        ws1 = wb.active
        ws1.title = "Báo cáo kinh doanh"
        ws1.views.sheetView[0].showGridLines = True
        
        ws1.merge_cells('A1:O1')
        title_cell = ws1['A1']
        title_cell.value = f"BÁO CÁO KINH DOANH & LỢI NHUẬN THÁNG {month:02d}/{year}"
        title_cell.font = font_title
        title_cell.alignment = align_center
        ws1.row_dimensions[1].height = 30
        
        ws1.merge_cells('A2:O2')
        sub_cell = ws1['A2']
        sub_cell.value = f"Ngày xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Đơn vị tiền tệ: VNĐ | Nguồn dữ liệu: Hệ thống Quản trị Vận hành Gido"
        sub_cell.font = font_subtitle
        sub_cell.alignment = align_center
        ws1.row_dimensions[2].height = 20
        
        achieve_val = (kpi['achievement_pct'] / 100.0) if kpi.get('achievement_pct') is not None else "Chưa có target"
        achieve_fmt = num_fmt_percent if kpi.get('achievement_pct') is not None else "@"
        
        kpi_cards = [
            ("DOANH THU BÁN (chưa VAT)", kpi['revenue'], num_fmt_currency),
            ("CƯỚC MUA VÀO (có VAT)", kpi['buy_cost'], num_fmt_currency),
            ("CHI PHÍ VẬN HÀNH", kpi['operating_cost'], num_fmt_currency),
            ("LỢI NHUẬN RÒNG", kpi['net_profit'], num_fmt_currency),
            ("TỶ SUẤT LN", kpi['profit_margin'] / 100.0, num_fmt_percent),
            ("TIẾN ĐỘ TARGET", achieve_val, achieve_fmt)
        ]
        
        ws1.row_dimensions[4].height = 18
        ws1.row_dimensions[5].height = 22
        
        col_start = 2
        for title, val, fmt in kpi_cards:
            col_end = col_start + 1
            ws1.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_end)
            ws1.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_end)
            
            c_top = ws1.cell(row=4, column=col_start, value=title)
            c_top.font = font_kpi_title
            c_top.alignment = align_center
            
            c_bot = ws1.cell(row=5, column=col_start, value=val)
            c_bot.font = font_kpi_val
            c_bot.alignment = align_center
            c_bot.number_format = fmt
            
            for r in [4, 5]:
                for c in [col_start, col_end]:
                    cell = ws1.cell(row=r, column=c)
                    cell.fill = fill_kpi
                    cell.border = thin_border
                    
            col_start += 2
            
        headers1 = [
            "STT", "Mã / Tên Lô", "Khách Hàng", "Công Ty", "Số Tờ Khai HQ", 
            "Ngày Bắt Đầu", "Ngày Hoàn Thành", "Doanh Thu Bán\n(chưa VAT)", 
            "Cước Mua Vào\n(có VAT)", "Chi Phí Vận Hành\n(thực tế)", 
            "Lợi Nhuận Gộp", "Lợi Nhuận Ròng", "Tỷ Suất LN", "Người Phụ Trách", "Trạng Thái"
        ]
        
        ws1.row_dimensions[7].height = 32
        for col_idx, h_text in enumerate(headers1, 1):
            cell = ws1.cell(row=7, column=col_idx, value=h_text.replace('\\n', '\n'))
            cell.font = font_header
            cell.fill = fill_navy
            cell.alignment = align_header
            cell.border = thin_border
            
        start_row1 = 8
        for i, lot in enumerate(lots, 1):
            curr_row = start_row1 + i - 1
            ws1.row_dimensions[curr_row].height = 22
            row_fill = fill_alt if i % 2 == 0 else PatternFill(fill_type=None)
            
            cust_name = lot.customer.name if lot.customer else "Khách vãng lai"
            assignee_name = lot.assignee.full_name if lot.assignee else "Chưa gán"
            start_d = lot.start_date.strftime('%d/%m/%Y') if lot.start_date else ""
            end_d = lot.end_date.strftime('%d/%m/%Y') if lot.end_date else ""
            
            status_map = {
                'pending': 'Chờ xử lý',
                'assigned': 'Đã giao việc',
                'in_progress': 'Đang nhập chi phí',
                'completed': 'Hoàn thành'
            }
            status_text = status_map.get(lot.status, lot.status)
            
            ws1.cell(row=curr_row, column=1, value=i).alignment = align_center
            ws1.cell(row=curr_row, column=2, value=lot.lot_label).alignment = align_center
            ws1.cell(row=curr_row, column=3, value=cust_name).alignment = align_left
            ws1.cell(row=curr_row, column=4, value=lot.company or "").alignment = align_left
            ws1.cell(row=curr_row, column=5, value=lot.customs_declaration or "").alignment = align_center
            ws1.cell(row=curr_row, column=6, value=start_d).alignment = align_center
            ws1.cell(row=curr_row, column=7, value=end_d).alignment = align_center
            
            c_rev = ws1.cell(row=curr_row, column=8, value=lot.total_sell_revenue)
            c_rev.number_format = num_fmt_currency
            c_rev.alignment = align_right
            
            c_buy = ws1.cell(row=curr_row, column=9, value=lot.total_buy_cost)
            c_buy.number_format = num_fmt_currency
            c_buy.alignment = align_right
            
            c_ops = ws1.cell(row=curr_row, column=10, value=lot.total_operating_cost)
            c_ops.number_format = num_fmt_currency
            c_ops.alignment = align_right
            
            c_gross = ws1.cell(row=curr_row, column=11, value=f"=H{curr_row}-I{curr_row}")
            c_gross.number_format = num_fmt_currency
            c_gross.alignment = align_right
            
            c_net = ws1.cell(row=curr_row, column=12, value=f"=H{curr_row}-I{curr_row}-J{curr_row}")
            c_net.number_format = num_fmt_currency
            c_net.alignment = align_right
            
            c_margin = ws1.cell(row=curr_row, column=13, value=f"=IF(H{curr_row}>0, L{curr_row}/H{curr_row}, 0)")
            c_margin.number_format = num_fmt_percent
            c_margin.alignment = align_right
            
            ws1.cell(row=curr_row, column=14, value=assignee_name).alignment = align_center
            ws1.cell(row=curr_row, column=15, value=status_text).alignment = align_center
            
            for c in range(1, 16):
                cell = ws1.cell(row=curr_row, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_fill.fill_type:
                    cell.fill = row_fill
                    
        last_data_row = start_row1 + len(lots) - 1
        
        if unresolved_cost_total > 0:
            # Row for Sales Lots subtotal
            sub_row = last_data_row + 1
            ws1.row_dimensions[sub_row].height = 24
            ws1.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=7)
            c_sub_label = ws1.cell(row=sub_row, column=1, value="TỔNG LÔ BÁN HÀNG")
            c_sub_label.font = font_total
            c_sub_label.alignment = align_center
            if len(lots) > 0:
                ws1.cell(row=sub_row, column=8, value=f"=SUM(H{start_row1}:H{last_data_row})")
                ws1.cell(row=sub_row, column=9, value=f"=SUM(I{start_row1}:I{last_data_row})")
                ws1.cell(row=sub_row, column=10, value=f"=SUM(J{start_row1}:J{last_data_row})")
                ws1.cell(row=sub_row, column=11, value=f"=SUM(K{start_row1}:K{last_data_row})")
                ws1.cell(row=sub_row, column=12, value=f"=SUM(L{start_row1}:L{last_data_row})")
                ws1.cell(row=sub_row, column=13, value=f"=IF(H{sub_row}>0, L{sub_row}/H{sub_row}, 0)")
            for col_idx in range(8, 13):
                cell = ws1.cell(row=sub_row, column=col_idx)
                cell.number_format = num_fmt_currency
                cell.alignment = align_right
                cell.font = font_total
            ws1.cell(row=sub_row, column=13).number_format = num_fmt_percent
            ws1.cell(row=sub_row, column=13).alignment = align_right
            ws1.cell(row=sub_row, column=13).font = font_total
            for c in range(1, 16):
                cell = ws1.cell(row=sub_row, column=c)
                cell.fill = fill_kpi
                cell.border = thin_border

            # Row for Unresolved CPVH
            unres_row = sub_row + 1
            ws1.row_dimensions[unres_row].height = 24
            ws1.merge_cells(start_row=unres_row, start_column=1, end_row=unres_row, end_column=7)
            c_unres_label = ws1.cell(row=unres_row, column=1, value=f"CHI PHÍ VẬN HÀNH CHƯA ĐỐI SOÁT ({len(unresolved_groups)} nhóm)")
            c_unres_label.font = font_total
            c_unres_label.alignment = align_center
            ws1.cell(row=unres_row, column=8, value=0).number_format = num_fmt_currency
            ws1.cell(row=unres_row, column=9, value=0).number_format = num_fmt_currency
            c_uops = ws1.cell(row=unres_row, column=10, value=unresolved_cost_total)
            c_uops.number_format = num_fmt_currency
            c_uops.alignment = align_right
            c_uops.font = font_total
            ws1.cell(row=unres_row, column=11, value=0).number_format = num_fmt_currency
            c_unet = ws1.cell(row=unres_row, column=12, value=-unresolved_cost_total)
            c_unet.number_format = num_fmt_currency
            c_unet.alignment = align_right
            c_unet.font = font_total
            ws1.cell(row=unres_row, column=13, value=0).number_format = num_fmt_percent
            for c in range(1, 16):
                cell = ws1.cell(row=unres_row, column=c)
                cell.fill = fill_alt
                cell.border = thin_border

            total_row = unres_row + 1
            ws1.row_dimensions[total_row].height = 26
            ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
            c_tot_label = ws1.cell(row=total_row, column=1, value="TỔNG CỘNG KỲ (BAO GỒM CPVH CHƯA GHÉP)")
            c_tot_label.font = font_total
            c_tot_label.alignment = align_center
            ws1.cell(row=total_row, column=8, value=f"=H{sub_row}")
            ws1.cell(row=total_row, column=9, value=f"=I{sub_row}")
            ws1.cell(row=total_row, column=10, value=f"=J{sub_row}+J{unres_row}")
            ws1.cell(row=total_row, column=11, value=f"=K{sub_row}")
            ws1.cell(row=total_row, column=12, value=f"=L{sub_row}+L{unres_row}")
            ws1.cell(row=total_row, column=13, value=f"=IF(H{total_row}>0, L{total_row}/H{total_row}, 0)")
        else:
            total_row = last_data_row + 1
            ws1.row_dimensions[total_row].height = 26
            ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
            c_tot_label = ws1.cell(row=total_row, column=1, value="TỔNG CỘNG")
            c_tot_label.font = font_total
            c_tot_label.alignment = align_center
            if len(lots) > 0:
                ws1.cell(row=total_row, column=8, value=f"=SUM(H{start_row1}:H{last_data_row})")
                ws1.cell(row=total_row, column=9, value=f"=SUM(I{start_row1}:I{last_data_row})")
                ws1.cell(row=total_row, column=10, value=f"=SUM(J{start_row1}:J{last_data_row})")
                ws1.cell(row=total_row, column=11, value=f"=SUM(K{start_row1}:K{last_data_row})")
                ws1.cell(row=total_row, column=12, value=f"=SUM(L{start_row1}:L{last_data_row})")
                ws1.cell(row=total_row, column=13, value=f"=IF(H{total_row}>0, L{total_row}/H{total_row}, 0)")
            else:
                for c in range(8, 14):
                    ws1.cell(row=total_row, column=c, value=0)

        for col_idx in range(8, 13):
            cell = ws1.cell(row=total_row, column=col_idx)
            cell.number_format = num_fmt_currency
            cell.alignment = align_right
            cell.font = font_total
            
        c_tot_margin = ws1.cell(row=total_row, column=13)
        c_tot_margin.number_format = num_fmt_percent
        c_tot_margin.alignment = align_right
        c_tot_margin.font = font_total
        
        for c in range(1, 16):
            cell = ws1.cell(row=total_row, column=c)
            cell.fill = fill_total
            cell.border = total_border

        # SHEET 2: CHI TIẾT CPVH
        ws2 = wb.create_sheet(title="Chi tiết CPVH")
        ws2.views.sheetView[0].showGridLines = True
        
        ws2.merge_cells('A1:R1')
        title_cell2 = ws2['A1']
        title_cell2.value = f"BẢNG CHI TIẾT CHI PHÍ VẬN HÀNH THÁNG {month:02d}/{year}"
        title_cell2.font = font_title
        title_cell2.alignment = align_center
        ws2.row_dimensions[1].height = 30
        
        headers2 = [
            "STT", "Mã Lô", "Khách Hàng", "Loại Chi Phí", "Nội Dung Chi Tiết", 
            "BKS Xe", "Số Lượng Xe", "Đơn Giá (VNĐ)", "Thành Tiền (VNĐ)", 
            "Loại Chứng Từ", "Ký Hiệu HĐ", "Số HĐ / Vé Xe", "Ngày Chứng Từ", 
            "MST NCC", "Nhà Cung Cấp", "Ghi Chú", "PIC Điền", "Phương Thức TT"
        ]
        
        ws2.row_dimensions[3].height = 30
        for col_idx, h_text in enumerate(headers2, 1):
            cell = ws2.cell(row=3, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_navy
            cell.alignment = align_header
            cell.border = thin_border
            
        all_period_lots = Lot.query.filter_by(month=month, year=year, is_deleted=False).all()
        lot_ids = [l.id for l in all_period_lots]
        costs = OperatingCost.query.filter(
            OperatingCost.lot_id.in_(lot_ids), OperatingCost.is_deleted.is_(False)
        ).order_by(OperatingCost.id.asc()).all() if lot_ids else []
        
        start_row2 = 4
        for idx, cost in enumerate(costs, 1):
            curr_row = start_row2 + idx - 1
            ws2.row_dimensions[curr_row].height = 20
            row_fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
            
            lot = cost.lot
            is_unres = bool(lot and lot.source_type == 'cpvh')
            cust_name = (lot.company or "Chưa đối soát") if is_unres else (lot.customer.name if (lot and lot.customer) else "Khách vãng lai")
            lot_name = f"[{lot.lot_label}] (Chưa đối soát)" if is_unres else (lot.lot_label if lot else f"Lô #{cost.lot_id}")
            doc_d = cost.document_date.strftime('%d/%m/%Y') if cost.document_date else ""
            filler_name = cost.filler.full_name if cost.filler else (cost.pic or "")
            
            ws2.cell(row=curr_row, column=1, value=idx).alignment = align_center
            ws2.cell(row=curr_row, column=2, value=lot_name).alignment = align_center
            ws2.cell(row=curr_row, column=3, value=cust_name).alignment = align_left
            ws2.cell(row=curr_row, column=4, value=cost.cost_type or "Chi phí khác").alignment = align_left
            ws2.cell(row=curr_row, column=5, value=cost.description).alignment = align_left
            ws2.cell(row=curr_row, column=6, value=cost.vehicle_plate or "").alignment = align_center
            ws2.cell(row=curr_row, column=7, value=cost.vehicle_count or 1).alignment = align_center
            
            c_uprice = ws2.cell(row=curr_row, column=8, value=cost.unit_price or 0)
            c_uprice.number_format = num_fmt_currency
            c_uprice.alignment = align_right
            
            c_total = ws2.cell(row=curr_row, column=9, value=cost.total_amount or 0)
            c_total.number_format = num_fmt_currency
            c_total.alignment = align_right
            
            ws2.cell(row=curr_row, column=10, value=cost.invoice_type or "").alignment = align_center
            ws2.cell(row=curr_row, column=11, value=cost.invoice_symbol or "").alignment = align_center
            ws2.cell(row=curr_row, column=12, value=cost.invoice_number or "").alignment = align_center
            ws2.cell(row=curr_row, column=13, value=doc_d).alignment = align_center
            ws2.cell(row=curr_row, column=14, value=cost.supplier_tax_code or "").alignment = align_center
            ws2.cell(row=curr_row, column=15, value=cost.supplier_name or "").alignment = align_left
            ws2.cell(row=curr_row, column=16, value=cost.note or "").alignment = align_left
            ws2.cell(row=curr_row, column=17, value=filler_name).alignment = align_center
            ws2.cell(row=curr_row, column=18, value=cost.payment_method or "Tiền mặt").alignment = align_center
            
            for c in range(1, 19):
                cell = ws2.cell(row=curr_row, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_fill.fill_type:
                    cell.fill = row_fill
                    
        last_cost_row = start_row2 + len(costs) - 1
        tot_cost_row = last_cost_row + 1
        ws2.row_dimensions[tot_cost_row].height = 24
        ws2.merge_cells(start_row=tot_cost_row, start_column=1, end_row=tot_cost_row, end_column=8)
        c_lbl = ws2.cell(row=tot_cost_row, column=1, value="TỔNG CHI PHÍ VẬN HÀNH")
        c_lbl.font = font_total
        c_lbl.alignment = align_center
        
        if len(costs) > 0:
            c_sum = ws2.cell(row=tot_cost_row, column=9, value=f"=SUM(I{start_row2}:I{last_cost_row})")
        else:
            c_sum = ws2.cell(row=tot_cost_row, column=9, value=0)
        c_sum.font = font_total
        c_sum.number_format = num_fmt_currency
        c_sum.alignment = align_right
        
        for c in range(1, 19):
            cell = ws2.cell(row=tot_cost_row, column=c)
            cell.fill = fill_total
            cell.border = total_border

        # SHEET 3: CƠ CẤU KHÁCH HÀNG
        ws3 = wb.create_sheet(title="Cơ cấu khách hàng")
        ws3.views.sheetView[0].showGridLines = True
        
        ws3.merge_cells('A1:E1')
        title_cell3 = ws3['A1']
        title_cell3.value = f"BÁO CÁO CƠ CẤU KHÁCH HÀNG THÁNG {month:02d}/{year}"
        title_cell3.font = font_title
        title_cell3.alignment = align_center
        ws3.row_dimensions[1].height = 30
        
        headers3 = ["STT", "Tên Khách Hàng", "Số Lô Hàng", "Tổng Doanh Thu (VNĐ)", "Tỷ Trọng Doanh Thu"]
        ws3.row_dimensions[3].height = 28
        for col_idx, h_text in enumerate(headers3, 1):
            cell = ws3.cell(row=3, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_navy
            cell.alignment = align_header
            cell.border = thin_border
            
        customers = CalculatorService.get_customer_breakdown(month, year)
        start_row3 = 4
        for idx, cust in enumerate(customers, 1):
            curr_row = start_row3 + idx - 1
            ws3.row_dimensions[curr_row].height = 20
            row_fill = fill_alt if idx % 2 == 0 else PatternFill(fill_type=None)
            
            ws3.cell(row=curr_row, column=1, value=idx).alignment = align_center
            ws3.cell(row=curr_row, column=2, value=cust['name']).alignment = align_left
            ws3.cell(row=curr_row, column=3, value=cust['lot_count']).alignment = align_center
            
            c_rev = ws3.cell(row=curr_row, column=4, value=cust['revenue'])
            c_rev.number_format = num_fmt_currency
            c_rev.alignment = align_right
            
            c_share = ws3.cell(row=curr_row, column=5, value=cust['share'] / 100.0)
            c_share.number_format = num_fmt_percent
            c_share.alignment = align_right
            
            for c in range(1, 6):
                cell = ws3.cell(row=curr_row, column=c)
                cell.font = font_data
                cell.border = thin_border
                if row_fill.fill_type:
                    cell.fill = row_fill
                    
        last_cust_row = start_row3 + len(customers) - 1
        tot_cust_row = last_cust_row + 1
        ws3.row_dimensions[tot_cust_row].height = 24
        
        ws3.merge_cells(start_row=tot_cust_row, start_column=1, end_row=tot_cust_row, end_column=2)
        c_lbl3 = ws3.cell(row=tot_cust_row, column=1, value="TỔNG CỘNG")
        c_lbl3.font = font_total
        c_lbl3.alignment = align_center
        
        if len(customers) > 0:
            ws3.cell(row=tot_cust_row, column=3, value=f"=SUM(C{start_row3}:C{last_cust_row})")
            ws3.cell(row=tot_cust_row, column=4, value=f"=SUM(D{start_row3}:D{last_cust_row})")
            ws3.cell(row=tot_cust_row, column=5, value=f"=SUM(E{start_row3}:E{last_cust_row})")
        else:
            ws3.cell(row=tot_cust_row, column=3, value=0)
            ws3.cell(row=tot_cust_row, column=4, value=0)
            ws3.cell(row=tot_cust_row, column=5, value=0)
            
        ws3.cell(row=tot_cust_row, column=3).alignment = align_center
        ws3.cell(row=tot_cust_row, column=3).font = font_total
        
        ws3.cell(row=tot_cust_row, column=4).number_format = num_fmt_currency
        ws3.cell(row=tot_cust_row, column=4).alignment = align_right
        ws3.cell(row=tot_cust_row, column=4).font = font_total
        
        ws3.cell(row=tot_cust_row, column=5).number_format = num_fmt_percent
        ws3.cell(row=tot_cust_row, column=5).alignment = align_right
        ws3.cell(row=tot_cust_row, column=5).font = font_total
        
        for c in range(1, 6):
            cell = ws3.cell(row=tot_cust_row, column=c)
            cell.fill = fill_total
            cell.border = total_border

        # AUTO-FIT COLUMN WIDTHS
        for sheet in [ws1, ws2, ws3]:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        if cell.row in [1, 2] and col_letter in ['A', 'B', 'C', 'D', 'E', 'F']:
                            continue
                        val_str = str(cell.value)
                        lines = val_str.split('\n')
                        for line in lines:
                            max_len = max(max_len, len(line))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
                
        ws1.column_dimensions['A'].width = 8
        ws1.column_dimensions['B'].width = 14
        ws1.column_dimensions['C'].width = 24
        ws1.column_dimensions['D'].width = 20
        ws1.column_dimensions['E'].width = 18
        ws1.column_dimensions['H'].width = 20
        ws1.column_dimensions['I'].width = 20
        ws1.column_dimensions['J'].width = 20
        ws1.column_dimensions['K'].width = 18
        ws1.column_dimensions['L'].width = 18
        ws1.column_dimensions['M'].width = 14
        ws1.column_dimensions['N'].width = 18
        ws1.column_dimensions['O'].width = 16
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_problematic_costs(lot_id: int = None, month: int = None, year: int = None) -> io.BytesIO:
        """
        Xuất file Excel danh sách các chi phí:
        1. Không có hóa đơn (no_invoice)
        2. Có hóa đơn nhưng không thanh toán được (unpayable_invoice: Phiếu chi, Phiếu thu, Vé xe...)
        Dùng để gửi email/báo cáo trình Sếp xem xét duyệt chi.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Chi phí cần xử lý"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Arial", size=14, bold=True, color="991B1B")
        font_subtitle = Font(name="Arial", size=9.5, italic=True, color="475569")
        font_header = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
        font_kpi_title = Font(name="Arial", size=8.5, bold=True, color="475569")
        font_kpi_val = Font(name="Arial", size=11, bold=True, color="0F172A")
        font_data = Font(name="Arial", size=9.0, color="1E293B")
        font_total = Font(name="Arial", size=10, bold=True, color="0F172A")

        fill_header = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        fill_alt = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
        fill_kpi = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_total = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_side = Side(style='thin', color='CBD5E1')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Side(style='double', color='991B1B')
        total_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)

        num_fmt_currency = '#,##0'

        # Query relevant costs
        costs_data = []
        scope_title = ""

        if lot_id:
            lot = Lot.query.get(lot_id)
            if lot:
                scope_title = f"LÔ HÀNG: {lot.lot_label} - {lot.display_customer_name} (Tháng {lot.month:02d}/{lot.year})"
                for c in lot.active_operating_costs:
                    if c.invoice_classification in ('no_invoice', 'unpayable_invoice'):
                        costs_data.append((lot, c))
        elif month and year:
            scope_title = f"KỲ BÁO CÁO: THÁNG {month:02d}/{year}"
            lots = Lot.query.filter_by(month=month, year=year, is_deleted=False).order_by(Lot.id.asc()).all()
            for lot in lots:
                for c in lot.active_operating_costs:
                    if c.invoice_classification in ('no_invoice', 'unpayable_invoice'):
                        costs_data.append((lot, c))
        else:
            scope_title = "TOÀN BỘ CÁC KỲ BÁO CÁO"
            lots = Lot.query.filter_by(is_deleted=False).order_by(Lot.year.desc(), Lot.month.desc(), Lot.id.asc()).all()
            for lot in lots:
                for c in lot.active_operating_costs:
                    if c.invoice_classification in ('no_invoice', 'unpayable_invoice'):
                        costs_data.append((lot, c))

        # Title Block
        ws.merge_cells('A1:R1')
        t_cell = ws['A1']
        t_cell.value = "DANH SÁCH HẠNG MỤC CHI PHÍ CẦN XỬ LÝ (KHÔNG HÓA ĐƠN & HÓA ĐƠN KHÔNG THANH TOÁN ĐƯỢC)"
        t_cell.font = font_title
        t_cell.alignment = align_center
        ws.row_dimensions[1].height = 28

        ws.merge_cells('A2:R2')
        sub_cell = ws['A2']
        sub_cell.value = f"{scope_title} | Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Đơn vị tiền tệ: VNĐ | Trình Lãnh đạo xem xét duyệt chi"
        sub_cell.font = font_subtitle
        sub_cell.alignment = align_center
        ws.row_dimensions[2].height = 18

        # KPI Summary Cards
        no_inv_items = [c for _, c in costs_data if c.invoice_classification == 'no_invoice']
        unpay_items = [c for _, c in costs_data if c.invoice_classification == 'unpayable_invoice']
        total_no_inv_amount = sum(c.total_amount or 0 for c in no_inv_items)
        total_unpay_amount = sum(c.total_amount or 0 for c in unpay_items)
        grand_total_amount = total_no_inv_amount + total_unpay_amount

        kpi_cards = [
            ("TỔNG HẠNG MỤC CẦN XỬ LÝ", f"{len(costs_data)} khoản", "@"),
            ("KHÔNG HÓA ĐƠN (SỐ TIỀN)", total_no_inv_amount, num_fmt_currency),
            ("HĐ KHÔNG THANH TOÁN ĐƯỢC", total_unpay_amount, num_fmt_currency),
            ("TỔNG TIỀN TRÌNH SẾP DUYỆT", grand_total_amount, num_fmt_currency)
        ]

        ws.row_dimensions[4].height = 18
        ws.row_dimensions[5].height = 22

        col_start = 2
        for title, val, fmt in kpi_cards:
            col_end = col_start + 3
            ws.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_end)
            ws.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_end)

            c_top = ws.cell(row=4, column=col_start, value=title)
            c_top.font = font_kpi_title
            c_top.alignment = align_center

            c_bot = ws.cell(row=5, column=col_start, value=val)
            c_bot.font = font_kpi_val
            c_bot.alignment = align_center
            c_bot.number_format = fmt

            for r in [4, 5]:
                for c in range(col_start, col_end + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = fill_kpi
                    cell.border = thin_border

            col_start += 4

        # Table Headers
        headers = [
            ("STT", align_center, 6),
            ("Lô Hàng", align_center, 14),
            ("Khách Hàng", align_left, 18),
            ("Phân Loại Vấn Đề", align_center, 22),
            ("Loại Chi Phí", align_center, 16),
            ("Nội Dung Chi Tiết", align_left, 32),
            ("Biển Số Xe", align_center, 12),
            ("SL Xe", align_center, 8),
            ("Đơn Giá", align_right, 14),
            ("Thành Tiền (CP)", align_right, 16),
            ("Loại Chứng Từ Hiện Có", align_center, 22),
            ("Ký Hiệu", align_center, 12),
            ("Số HĐ / Phiếu", align_center, 14),
            ("Ngày Chứng Từ", align_center, 14),
            ("Nhà Cung Cấp", align_left, 24),
            ("MST NCC", align_center, 14),
            ("PIC Phụ Trách", align_left, 16),
            ("Ghi Chú", align_left, 22)
        ]

        ws.row_dimensions[7].height = 26
        for col_idx, (hdr, alignment, width) in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col_idx, value=hdr)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_header
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Data Rows
        current_row = 8
        for idx, (lot, cost) in enumerate(costs_data, 1):
            is_alt = (idx % 2 == 0)
            row_fill = fill_alt if is_alt else PatternFill(fill_type=None)

            prob_label = "Không có HĐ" if cost.invoice_classification == 'no_invoice' else "HĐ không thanh toán được"
            doc_date_str = cost.document_date.strftime('%d/%m/%Y') if cost.document_date else ''

            row_data = [
                (idx, align_center, "@"),
                (lot.lot_label if lot else '', align_center, "@"),
                (lot.display_customer_name if lot else '', align_left, "@"),
                (prob_label, align_center, "@"),
                (cost.display_cost_type, align_center, "@"),
                (cost.display_description, align_left, "@"),
                (cost.vehicle_plate or '-', align_center, "@"),
                (cost.vehicle_count or 1, align_center, "#,##0"),
                (cost.unit_price or 0, align_right, num_fmt_currency),
                (cost.total_amount or 0, align_right, num_fmt_currency),
                (cost.invoice_type or 'Chưa có chứng từ', align_center, "@"),
                (cost.invoice_symbol or '-', align_center, "@"),
                (cost.invoice_number or '-', align_center, "@"),
                (doc_date_str, align_center, "@"),
                (cost.supplier_name or '-', align_left, "@"),
                (cost.supplier_tax_code or '-', align_center, "@"),
                (cost.pic or (cost.filler.full_name if cost.filler else '-'), align_left, "@"),
                (cost.note or '', align_left, "@")
            ]

            ws.row_dimensions[current_row].height = 20
            for col_idx, (val, alignment, fmt) in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_data
                cell.alignment = alignment
                cell.border = thin_border
                cell.number_format = fmt
                if is_alt:
                    cell.fill = row_fill
                # Highlight problem label
                if col_idx == 4:
                    if cost.invoice_classification == 'no_invoice':
                        cell.font = Font(name="Arial", size=9.0, bold=True, color="DC2626")
                    else:
                        cell.font = Font(name="Arial", size=9.0, bold=True, color="D97706")

            current_row += 1

        # Total Row
        ws.row_dimensions[current_row].height = 24
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        tot_label = ws.cell(row=current_row, column=1, value="TỔNG CỘNG CHI PHÍ CẦN XỬ LÝ:")
        tot_label.font = font_total
        tot_label.alignment = align_right

        tot_val = ws.cell(row=current_row, column=10, value=grand_total_amount)
        tot_val.font = font_total
        tot_val.alignment = align_right
        tot_val.number_format = num_fmt_currency

        for c in range(1, 19):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = fill_total
            cell.border = total_border

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

